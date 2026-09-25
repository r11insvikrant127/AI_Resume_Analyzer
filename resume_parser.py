import re
import io

from pypdf import PdfReader
from docx import Document


def extract_text_from_pdf(uploaded_file):
    try:
        reader = PdfReader(uploaded_file)
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
        return "\n".join(pages)
    except Exception as e:
        raise Exception(f"Unable to read PDF: {e}")


def extract_text_from_docx(uploaded_file):
    try:
        # Streamlit UploadedFile is file-like; python-docx needs a stream
        file_bytes = uploaded_file.read()
        document = Document(io.BytesIO(file_bytes))

        parts = []

        # Paragraphs
        for paragraph in document.paragraphs:
            if paragraph.text:
                parts.append(paragraph.text)

        # Tables (resumes often use tables for layout)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        parts.append(cell.text)

        return "\n".join(parts)
    except Exception as e:
        raise Exception(f"Unable to read DOCX: {e}")


def extract_resume_text(uploaded_file):
    """
    Extract text from a PDF or DOCX resume.

    Detection is based on the file extension.
    """
    name = (getattr(uploaded_file, "name", "") or "").lower()

    if name.endswith(".pdf"):
        return extract_text_from_pdf(uploaded_file)

    if name.endswith(".docx"):
        return extract_text_from_docx(uploaded_file)

    raise Exception(
        "Unsupported file type. Please upload a PDF or DOCX."
    )


def clean_resume_text(text):
    if not text:
        return ""

    text = re.sub(r"\s+", " ", text)
    return text.strip()